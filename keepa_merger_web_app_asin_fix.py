
import streamlit as st
import pandas as pd
from io import BytesIO
import openpyxl
from openpyxl.styles import PatternFill

st.set_page_config(page_title="Keepa Merger Tool", layout="centered")

st.title("📦 Keepa USA/UK Product Merger Tool")
st.markdown("Upload your **USA Seller ASIN export** and the **UK Product Finder export** from Keepa to generate FILE-3.")

file1 = st.file_uploader("🇺🇸 Upload Keepa USA Export (Seller ASIN List)", type="csv")
file2 = st.file_uploader("🇬🇧 Upload Keepa UK Export (Product Finder Results)", type="csv")

required_cols_file1 = ['ASIN', 'Title', 'Buy Box 🚚: Current']
required_cols_file2 = ['ASIN', 'Buy Box 🚚: Current', 'Buy Box Seller']

def validate_columns(df, required):
    return all(col in df.columns for col in required)

if file1 and file2:
    try:
        df_usa = pd.read_csv(file1, dtype={'ASIN': str})
        df_uk = pd.read_csv(file2, dtype={'ASIN': str})

        df_usa.columns = df_usa.columns.str.strip()
        df_uk.columns = df_uk.columns.str.strip()

        valid_usa = validate_columns(df_usa, required_cols_file1)
        valid_uk = validate_columns(df_uk, required_cols_file2)

        if not valid_usa:
            st.error(f"❌ USA file is missing one or more required columns: {required_cols_file1}")
        elif not valid_uk:
            st.error(f"❌ UK file is missing one or more required columns: {required_cols_file2}")
        else:
            df_usa_clean = df_usa[required_cols_file1].copy()
            df_usa_clean.rename(columns={
                'Title': 'USA Title',
                'Buy Box 🚚: Current': 'USA BuyBox Price (USD)'
            }, inplace=True)
            df_usa_clean['Amazon.com Link'] = 'https://www.amazon.com/dp/' + df_usa_clean['ASIN']

            df_uk_clean = df_uk[required_cols_file2].copy()
            df_uk_clean.rename(columns={
                'Buy Box 🚚: Current': 'UK BuyBox Price (GBP)',
                'Buy Box Seller': 'UK Seller'
            }, inplace=True)
            df_uk_clean['Amazon.co.uk Link'] = 'https://www.amazon.co.uk/dp/' + df_uk_clean['ASIN']

            df = pd.merge(df_usa_clean, df_uk_clean, on='ASIN', how='inner')
            df.insert(1, 'Brand', df['USA Title'].str.split().str[0])
            df_final = df[['ASIN', 'Brand', 'USA Title', 'Amazon.com Link', 'Amazon.co.uk Link',
                           'USA BuyBox Price (USD)', 'UK BuyBox Price (GBP)', 'UK Seller']]

            st.success("✅ Files validated and merged successfully!")
            st.dataframe(df_final)

            # Save Excel (no index)
            excel_output = BytesIO()
            with pd.ExcelWriter(excel_output, engine='openpyxl') as writer:
                df_final.to_excel(writer, index=False, sheet_name='FILE-3')
                wb = writer.book
                ws = writer.sheets['FILE-3']
                yellow_fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=6)
                    try:
                        if cell.value and float(cell.value) < 15:
                            cell.fill = yellow_fill
                    except:
                        continue

            # Save CSV (no index)
            csv_output = df_final.to_csv(index=False).encode('utf-8')

            st.download_button(
                label="📥 Download FILE-3_Final.xlsx",
                data=excel_output.getvalue(),
                file_name="FILE-3_Final.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            st.download_button(
                label="📥 Download FILE-3_Final.csv",
                data=csv_output,
                file_name="FILE-3_Final.csv",
                mime="text/csv"
            )

    except Exception as e:
        st.error(f"❌ Error processing files: {e}")
